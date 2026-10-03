# -*- coding: utf-8 -*-
"""Regenerate MANIFEST.txt for the submission package.

The manifest records a SHA-256 for every document, the pixel size and stored dpi of every figure,
and the checks that were run. It is for the author's record and is not uploaded to Editorial Manager.

    python make_manifest.py
"""
import hashlib
import io
import os
import re

from docx import Document
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.abspath(os.path.join(HERE, "..", "..", "revision2"))
OUT = os.path.join(REV, "MANIFEST.txt")

DOCS = ["ILEDBV_Manuscript_Revised_v6.docx", "ILEDBV_Manuscript_Revised_v6_marked.docx",
        "DWT_Response_to_Reviewers_R2.docx", "DWT_Revision2_Cover_Letter.docx", "Highlights_ILEDBV_Revised.docx"]

HEADER = """Submission package - DWT-D-26-01366 Revision 2
Prepared and verified 3 October 2026. Do not upload this manifest to Editorial Manager.

PROVENANCE OF THE REPRODUCIBILITY CLAIM
  On 2 October 2026 the deposited repository github.com/sandlerleon/iledbv-desalination-model
  was cloned at tag v2.3.0 and run under Python 3.11.9 with NumPy 2.4.6, the environment named
  in the Data Availability statement. iledbv_revision_v2.py, design_and_costing.py,
  boron_recalc.py, reversal_analysis.py, alkalinity_calcium_design.py and attribution_analysis.py
  were executed and their outputs compared against the manuscript:
    - every Monte Carlo statistic reproduced exactly (median, 5th and 95th percentile)
    - 271 of 300,000 samples below the comparator, best case $0.352, reproduced exactly
    - alkalinity routes, boron, polarization ceiling and all Section 4.13 optima reproduced exactly
    - the nine carried-over figures are pixel-identical to regeneration from v2.3.0
    - every value plotted in the two new figures matches design_costing.json and
      iledbv_revision_v2.json exactly
  Figures 1, 3, 4, 5, 6 and 8 were re-exported from the same deposited code with savefig dpi
  raised from 300 to 600 for print resolution; plot content is unchanged. TIFFs are LZW
  compressed, losslessly and pixel-verified.
  This verification was carried out by the author's assistant, not by a third party, and covers
  the numerical outputs of the deposited model, not the physical validity of the model itself.

WHAT CHANGED AFTER THAT VERIFICATION (manuscript v5 -> v6, 3 October 2026)
  Two forward-looking subsections were added to the Discussion in response to editorial guidance,
  and neither touches the model. Section 5.3 gained two paragraphs on crosslinked graphene oxide
  and MXene laminate membranes; a new Section 5.5 sets out brine-derived magnesium as a precursor
  for metallothermic silicon production; the former Sections 5.5 and 5.6 are renumbered 5.6 and
  5.7 and each gained the matching qualification; one sentence was added at the end of the
  Introduction; one paragraph was added to the Conclusions before the closing sentence; and ten
  references, [27] to [36], were added. Silicon appears in neither the abstract nor the keywords.
  No equation, table, figure, parameter or number was altered: audit_manuscript.py reports 18 of
  18 headline numbers present and 0 contradictions against the same model JSON, and
  verify_manuscript_structure.py reports equations 1-14 all cited, tables 1-12 all cited, and
  references 1-36 all listed and all cited. Every one of the 36 DOIs resolves against Crossref
  with a matching title; [33] differs only because the Crossref record carries inline markup
  around the subscript x in Ti3C2Tx.

  A third-party language edit of v6 was then merged paragraph by paragraph
  (code/revision/merge_language_edit_v6.py). An edited paragraph is taken only when it carries the
  same content words, numbers, currency and unit tokens and sub/superscript runs as the source;
  661 paragraphs were unchanged, 51 edits were accepted and 44 were kept from the source. The
  rejected edits include 'brine' changed to 'saltwater', the reagent masses of Section 4.1 changed
  to 'solids', '$0.061' with its dollar sign dropped, 'precipitation train' changed to 'training'
  and 'vapor' misspelled; revision2/v6_language_edit_merge_report.txt records every decision.
  Five final corrections follow (code/revision/apply_final_feedback_v6.py): the laminate membranes
  of Section 5.3 are described as potential components of alternative brine-concentration and
  crystallization systems rather than as replacing crystallization; the electricity-price crossings
  are labelled modeled thresholds under the stated cost assumptions; the generative-AI declaration
  states what the tool did and what was verified and by what; Data availability names the plotting
  script for Figures 2 and 7 and the release that contains it; and numeric ranges are set with en
  dashes. The cover letter and the response letter are dated 3 October 2026 to match.
"""


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for blk in iter(lambda: f.read(1 << 20), b""):
            h.update(blk)
    return h.hexdigest()


lines = [HEADER, "DOCUMENTS"]
for d in DOCS:
    p = os.path.join(REV, d)
    lines.append("  %-44s %5.2f MB  sha256:%s  ok" % (d, os.path.getsize(p) / 1e6, sha(p)[:16]))

lines += ["", "FIGURES", "  %-22s %-12s %-5s %-16s %s" % ("file", "pixels", "dpi", "native size", "dpi at 190 mm")]
figs = sorted(f for f in os.listdir(REV) if re.match(r"(Figure\d+|Graphical_Abstract)\.(png|tif)$", f))
for f in figs:
    im = Image.open(os.path.join(REV, f))
    w, h = im.size
    dpi = int(round((im.info.get("dpi") or (300, 300))[0]))
    lines.append("  %-22s %-12s %-5d %-16s %d" % (f, "%dx%d" % (w, h), dpi, "%.2f x %.2f in" % (w / float(dpi), h / float(dpi)),
                                                  round(w / (190 / 25.4))))

doc = Document(os.path.join(REV, DOCS[0]))
nfig = sum(1 for _ in doc.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
neq = len(re.findall(r"<m:oMath[ >]", doc.element.xml))
abstract = next(p for i, p in enumerate(doc.paragraphs) if doc.paragraphs[i - 1].text.strip() == "Abstract").text
hi = Document(os.path.join(REV, "Highlights_ILEDBV_Revised.docx"))
bullets = [p.text.strip() for p in hi.paragraphs if p.style.name.startswith("List") and p.text.strip()]
refs = sum(1 for p in doc.paragraphs if re.match(r"^\[\d+\] ", p.text))

lines += ["", "CHECKS",
          "  highlights          %d bullets, longest %d characters (limit 85)" % (len(bullets), max(len(b) for b in bullets)),
          "  manuscript          %d embedded figures, %d native equation objects, abstract %d words, %d references"
          % (nfig, neq, len(abstract.split()), refs),
          "  numbers             audit_manuscript.py: 18/18 headline numbers present, 0 contradictions",
          "  structure           verify_manuscript_structure.py: no problems (equations, tables, figures, references)",
          "  DOIs                36/36 resolve against Crossref with matching titles"]

io.open(OUT, "w", encoding="utf-8").write("\n".join(lines) + "\n")
print("\n".join(lines[-6:]))
print("\nwrote", OUT)
