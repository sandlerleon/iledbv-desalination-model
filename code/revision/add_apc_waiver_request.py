# -*- coding: utf-8 -*-
"""Add an article-publishing-charge waiver request to the DWT Revision 2 cover letter.

Desalination and Water Treatment has been fully open access since 1 January 2024, with an article
publishing charge of up to USD 1,400 and no subscription route. The author is in the United States,
which is not covered by the Research4Life waiver, so the only available route is Elsevier's
discretionary waiver for genuine need. That is decided after acceptance, at the rights and access
step, but the editor is told now so that it is not raised late.

    python add_apc_waiver_request.py
"""
import copy
import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.abspath(os.path.join(HERE, "..", "..", "revision2"))
F = os.path.join(REV, "DWT_Revision2_Cover_Letter.docx")

LABEL = "Request for an article publishing charge waiver. "
BODY = (
    "I am an independent researcher with no institutional affiliation, no grant and no other source of funds for publication "
    "charges, and I am not able to pay the article publishing charge. I would be grateful if a full waiver could be considered "
    "under Elsevier's policy of granting waivers in cases of genuine need. This work received no funding of any kind, as the "
    "Funding declaration in the manuscript records, and the computation reported in it was carried out on personal equipment; the "
    "model code and the manuscript are already openly deposited at the DOIs given in the Data availability statement, so the work "
    "is freely available irrespective of the outcome of this request. I understand that the decision is made after acceptance, at "
    "the rights and access stage, and I raise it here only so that it does not arise late. I will complete any form or provide any "
    "statement the journal requires."
)

doc = Document(F)
anchor = next(i for i, p in enumerate(doc.paragraphs)
              if p.text.strip().startswith("The manuscript is original"))
src = doc.paragraphs[anchor]
new_p = copy.deepcopy(src._p)
src._p.addprevious(new_p)
from docx.text.paragraph import Paragraph  # noqa: E402

par = Paragraph(new_p, src._parent)
tmpl = par.runs[0]
for r in list(par.runs)[1:]:
    r._r.getparent().remove(r._r)
tmpl.text = LABEL
tmpl.bold = True
run = par.add_run(BODY)
run.bold = False
run.italic = tmpl.italic
run.font.size = tmpl.font.size
run.font.name = tmpl.font.name
doc.save(F)

chk = Document(F)
txt = "\n".join(p.text for p in chk.paragraphs)
assert txt.count(LABEL) == 1 and "full waiver" in txt
print("waiver request inserted before the originality paragraph")
for p in chk.paragraphs:
    if p.text.startswith(LABEL):
        print("\n" + p.text)
