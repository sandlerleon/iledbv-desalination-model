# -*- coding: utf-8 -*-
"""Bring DWT_Response_to_Reviewers.docx into line with manuscript v4.

The letter was written against an earlier numbering. This corrects every cross-
reference that no longer matches (tables, sections, equation and figure counts,
the release tag) and adds a paragraph to the summary describing what changed in
this version. Each replacement is asserted, so a silent miss cannot happen.

    python update_response_letter_v4.py

It edits the letter in place and is meant to be run once, on the letter as it
stood for v3; a second run stops at the first assertion because the old
wording is gone.
"""
import copy
import io
import json
import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
LETTER = os.path.join(REPO, "manuscript", "DWT_Response_to_Reviewers.docx")
AC = json.load(io.open(os.path.join(HERE, "alkalinity_calcium_results.json"), encoding="utf-8"))
RV = json.load(io.open(os.path.join(HERE, "reversal_results.json"), encoding="utf-8"))

FIXES = [
    ("Thirteen equations are now numbered consecutively",
     "Fourteen equations are now numbered consecutively"),
    ("and Table 4 gives the OPEX composition.", "and Table 5 gives the OPEX composition."),
    ("Table 4 now gives a nine-line OPEX breakdown", "Table 5 now gives a nine-line OPEX breakdown"),
    ("Lithium remains excluded from the baseline and this is stated in Section 5.4.",
     "Lithium remains excluded from the baseline and this is stated in Section 5.5."),
    ("replaced by the boundary comparison in Table 5.", "replaced by the boundary comparison in Table 6."),
    ("all 25 entries are cited and none is orphaned", "all 26 entries are cited and none is orphaned"),
    ("Table 5 sets out the boundary element by element", "Table 6 sets out the boundary element by element"),
    ("release tag v2.0.1", "release tag v2.1.0"),
    ("A Monte Carlo over eleven uncertain parameters", "A Monte Carlo over thirteen uncertain parameters"),
    ("(Table 7, Figure 5)", "(Tables 7 and 8, Figure 5)"),
    ("Section 5.5 proposes three specific bench-scale measurements",
     "Section 5.6 proposes three specific bench-scale measurements"),
    ("Table 3 gives the full design basis", "Table 4 gives the full design basis"),
    ("Discussed in Sections 4.3 and 5.4", "Discussed in Sections 4.3 and 5.5"),
    ("numbered (1)-(13); all seven tables are numbered", "numbered (1)-(14); all twelve tables are numbered"),
    ("seven captioned tables and six figures", "twelve captioned tables and nine figures"),
    ("Equations (1)-(13) are numbered and defined", "Equations (1)-(14) are numbered and defined"),
]


def set_text(par, text):
    runs = par.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def main():
    doc = Document(LETTER)
    for old, new in FIXES:
        hits = [p for p in doc.paragraphs if old in p.text]
        assert len(hits) == 1, "%r found %d times" % (old, len(hits))
        set_text(hits[0], hits[0].text.replace(old, new))

    base = AC["optimisation"]["model base case"]
    ceil = AC["price_ceiling_usd_per_kwh"]["reversible_limit"]
    carbon = {c["kg_per_kwh"]: c for c in RV["carbon"]}
    added = (
        "This version makes three further changes. First, Section 4.13 treats alkalinity "
        "supply and calcium closure as one design problem rather than route by route: a "
        "linear program, Eq. (14), chooses among lime, purchased NaOH and on-site "
        "electrochemical base, soda ash and dosed CO2, and a gypsum step that removes the "
        "calcium lime imports, with the calcium balance imposed as a constraint. Closing the "
        "balance costs $%.2f m-3 at the base case with the gypsum step and $%.2f without it, "
        "and above $%.3f kWh-1 no on-site base can close the loop against the comparator, "
        "because the reversible limit of water dissociation binds (Tables 11-12, Figure 9). "
        "Section 5.4 sets out what an alkalinity pathway would have to achieve on that basis. "
        "The paper is positioned as quantifying the conditions for closure, not as reporting "
        "a working process, and the abstract's electrochemical sentences now state that the "
        "result is modeled and exclude electrochemical capital and coproduct handling. "
        "Second, a double count was found and corrected in Section 4.12: the dosed CO2 of "
        "route D had been credited both as a reagent and again as carbon fixed in the "
        "product. On solar-dominated supply the route emits %+.2f kg CO2 m-3, not the "
        "+0.13 previously reported, and %+.1f on the Saudi grid mix rather than +20.3; the "
        "conclusion that route D requires low-carbon power is unchanged. Third, a structural "
        "check found that Figures 2-6 and Tables 5 and 8 were submitted but never cited in "
        "the text, and that Figure 5 had been drawn from a separate 200,000-sample run; all "
        "are now cited in order and Figure 5 is redrawn from the same 300,000-sample draw "
        "that produces Table 8."
        % (base["closure_cost_usd_per_m3"], base["closure_cost_without_gypsum"], ceil,
           carbon[0.05]["co2_D"], carbon[0.55]["co2_D"]))
    anchor = next(p for p in doc.paragraphs if p.text.startswith("One correction follows from that work."))
    new_p = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    set_text(Paragraph(new_p, anchor._parent), added)
    doc.save(LETTER)
    print("updated", LETTER)


if __name__ == "__main__":
    main()
