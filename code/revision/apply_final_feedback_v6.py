# -*- coding: utf-8 -*-
"""Five final corrections to manuscript v6 before resubmission, from the last round of feedback.

  1. Section 5.3 said the laminate membranes would be used "in place of conventional crystallization",
     which could be read as claiming they remove the need to crystallize. They are now described as
     potential components of alternative brine-concentration and crystallization systems, and the
     distinction is stated explicitly.
  2. The electricity-price thresholds are labelled modeled thresholds under the stated cost
     assumptions, because the capital cost of the electrochemical unit and its coproduct handling are
     outside the boundary.
  3. The generative-AI declaration is rewritten to describe what the tool actually did and what the
     author actually verified.
  4. Data availability now names the plotting script and the release that contains it: the script was
     added after v2.3.0, which is the tag cited for the model code, so citing only v2.3.0 for it was
     wrong.
  5. Numeric ranges written with hyphens are set with en dashes.

No model, parameter, table, figure or number changes.

    python apply_final_feedback_v6.py
"""
import io
import os
import re

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.abspath(os.path.join(HERE, "..", "..", "revision2"))
DOC = os.path.join(REV, "ILEDBV_Manuscript_Revised_v6.docx")

AI_DECL = (
    "During the preparation of this work the author used Claude (Anthropic) as a drafting and analysis aid: to draft and edit "
    "manuscript text, to write the Python that implements the model and the scripts that check it, to typeset the equations and "
    "tables, and to resolve each cited reference against its Crossref record. The author specified the architecture, the system "
    "boundary, the governing equations, the parameter ranges and their sources, and every assumption; directed and reviewed the "
    "analysis; inspected the model outputs and the figures; and decided the conclusions, including the finding that the original "
    "economic claim does not survive. Verification was performed by the deposited scripts rather than by hand: audit_manuscript.py "
    "rechecks the built document against the model JSON, and verify_manuscript_structure.py checks the equation, table, figure and "
    "reference numbering and re-resolves every DOI. That is a check of internal consistency and of the numbers as computed, not an "
    "independent check of the physical validity of the model. No AI tool is listed as an author, none generated or altered data, "
    "and the author takes full responsibility for the content of the publication."
)

REPL = [
    # 1. membrane terminology
    ("the use of two-dimensional laminate membranes in place of conventional crystallization.",
     "the use of two-dimensional laminate membranes as potential components of alternative brine-concentration and "
     "crystallization systems."),
    ("Neither material can presently be assumed to outperform mechanical vapor compression under the conditions modeled here.",
     "Membrane separation and complete crystallization are not interchangeable: a laminate membrane could at best concentrate the "
     "brine further and reduce the duty passed to the crystallizer, not remove the need to crystallize the residual salts. Neither "
     "material can presently be assumed to outperform mechanical vapor compression under the conditions modeled here."),
    ("Advanced GO- and MXene-based membrane configurations are identified in Section 5.3 as potential alternatives for future "
     "investigations,",
     "Advanced GO- and MXene-based membrane configurations are identified in Section 5.3 as potential components of alternative "
     "brine-concentration and crystallization systems for future investigations,"),
    # 2. modeled thresholds
    ("Solving for the price at which route D reaches the conventional comparator gives $0.033 kWh−1 (Fig. 10A).",
     "Solving for the price at which route D reaches the conventional comparator gives $0.033 kWh−1 (Fig. 10A). Both crossings "
     "are modeled thresholds under the cost assumptions of Tables 5 and 6 rather than demonstrated commercial break-even prices: "
     "the capital cost of the electrochemical unit and the handling of its chlorine and acid coproducts lie outside the boundary "
     "(Section 5.6)."),
    ("route D reaches the conventional comparator below $0.033 kWh−1, whereas route A remains $0.34 m−3 above it even at "
     "zero electricity cost.",
     "route D reaches the conventional comparator below a modeled electricity price of $0.033 kWh−1, whereas route A remains "
     "$0.34 m−3 above it even at zero electricity cost. That figure is a threshold computed under the stated cost assumptions, "
     "not a demonstrated commercial break-even price."),
    ("With on-site base, the closed loop reaches the comparator only if the base is generated at better than about 19 mol OH− "
     "per kWh from electricity at or below the $0.048 kWh−1 industrial tariff,",
     "With on-site base, and again under the stated cost assumptions, the closed loop reaches the comparator only if the base is "
     "generated at better than about 19 mol OH− per kWh from electricity at or below the $0.048 kWh−1 industrial tariff,"),
    # 4. data availability
    ("Figures 2 and 7 were added in the present revision and are plotted directly from Tables 1, 5 and 8 by a separate plotting "
     "script supplied with this submission for deposit alongside the model code; they introduce no new model output, and no number "
     "in this paper changed in this revision.",
     "Figures 2 and 7 were added in the present revision and are plotted directly from Tables 1, 5 and 8 by "
     "revision2/make_new_figures_v5.py, which is deposited in the same repository from release tag v2.5.0 onward; they introduce "
     "no new model output, and no number in this paper changed in this revision. The model code itself is unchanged since v2.3.0, "
     "which is the tag cited above for the results."),
]
DASHFIX = re.compile(r"(?<=\d)-(?=\d)")


def runs_of(doc):
    for p in doc.element.body.iter(qn("w:p")):
        yield p


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def replace_in_paragraph(p, old, new):
    """Replace `old` with `new` across run boundaries; the replacement takes the formatting of the run
    in which the match starts, so sub- and superscripts elsewhere in the paragraph are untouched."""
    ts = [t for t in p.iter(qn("w:t"))]
    full = "".join(t.text or "" for t in ts)
    start = full.find(old)
    if start < 0:
        return False
    end = start + len(old)
    pos, first = 0, True
    for t in ts:
        txt = t.text or ""
        a, b = pos, pos + len(txt)
        pos = b
        if b <= start or a >= end:
            continue
        lo, hi = max(start, a) - a, min(end, b) - a
        if first:
            t.text = txt[:lo] + new + txt[hi:]
            first = False
        else:
            t.text = txt[:lo] + txt[hi:]
    return True


doc = Document(DOC)
done = {old: 0 for old, _ in REPL}
for p in runs_of(doc):
    for old, new in REPL:
        if replace_in_paragraph(p, old, new):
            done[old] += 1

missing = [o[:60] for o, n in done.items() if n != 1]
assert not missing, "replacement not applied exactly once: %s" % missing

# 3. AI declaration
hit = 0
for p in doc.paragraphs:
    if p.text.startswith("During the preparation of this work") and "Claude" in p.text:
        for r in p.runs[1:]:
            r.text = ""
        p.runs[0].text = AI_DECL
        hit += 1
assert hit == 1, "AI declaration paragraph not found exactly once"

# 5. en dashes in numeric ranges, outside equations and reference entries
nd = 0
for p in doc.element.body.iter(qn("w:p")):
    if re.match(r"^\[\d+\] ", ptext(p).strip()):
        continue
    for t in p.iter(qn("w:t")):
        if t.text and DASHFIX.search(t.text):
            t.text, k = DASHFIX.subn("–", t.text)
            nd += k
doc.save(DOC)

chk = Document(DOC)
txt = "\n".join(p.text for p in chk.paragraphs)
assert "in place of conventional crystallization" not in txt
assert "modeled thresholds under the cost assumptions" in txt
assert "make_new_figures_v5.py" in txt
assert len(chk.paragraphs[7].text.split()) <= 250, "abstract exceeded 250 words"
print("replacements applied: %d | AI declaration rewritten | %d numeric ranges set with en dashes" % (len(REPL), nd))
print("abstract words:", len(chk.paragraphs[7].text.split()))
