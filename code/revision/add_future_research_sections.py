# -*- coding: utf-8 -*-
"""Add the two bounded future-research directions to the Revision 2 manuscript.

Editorial guidance received after the Revision 2 package was assembled asked for two research
directions to be stated without changing any modeled result: advanced two-dimensional laminate
membranes as an alternative to the crystallization stage, and brine-derived magnesium as a
precursor for metallothermic silicon production. Both are written as proposals, both are excluded
from the mass, energy, carbon and economic balances, and neither appears in the abstract or the
keywords.

    v5  ->  v6      python add_future_research_sections.py

What changes
  Introduction  one sentence placing the unmodeled extensions outside the system boundary
  5.3           two paragraphs on crosslinked GO and MXene laminates, with the comparison they
                would have to pass against the MVC reference
  5.5 (new)     brine-derived magnesium for metallothermic silicon production
  5.6 / 5.7     former 5.5 Limitations and 5.6 Proposed validation, renumbered, each extended
  6             one forward-looking paragraph before the closing sentence
  References    ten new entries, every one resolved against its Crossref record by harvest_refs

No table, figure, equation or number from the model is touched.
"""
import copy
import io
import json
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REV2 = os.path.abspath(os.path.join(HERE, "..", "..", "revision2"))
SRC = os.path.join(REV2, "ILEDBV_Manuscript_Revised_v5.docx")
OUT = os.path.join(REV2, "ILEDBV_Manuscript_Revised_v6.docx")

# ------------------------------------------------------------------ new references
# key -> (number assigned, formatted entry). Each DOI was resolved against Crossref
# (_refs_r3_raw.json holds the records) and the entry is formatted from that record.
NEW_REFS = [
    ("abraham2017", "J. Abraham, K.S. Vasu, C.D. Williams, K. Gopinadhan, Y. Su, C.T. Cherian, et al., Tunable sieving of ions "
     "using graphene oxide membranes, Nature Nanotechnology 12 (2017) 546-550. https://doi.org/10.1038/nnano.2017.21"),
    ("ding2020", "L. Ding, L. Li, Y. Liu, Y. Wu, Z. Lu, J. Deng, et al., Effective ion sieving with Ti{{s|3}}C{{s|2}}T{{s|x}} MXene "
     "membranes for production of drinking water from seawater, Nature Sustainability 3 (2020) 296-302. "
     "https://doi.org/10.1038/s41893-020-0474-0"),
    ("gao2009", "F. Gao, Z. Nie, Z. Wang, X. Gong, T. Zuo, Life cycle assessment of primary magnesium production using the Pidgeon "
     "process in China, The International Journal of Life Cycle Assessment 14 (2009) 480-489. "
     "https://doi.org/10.1007/s11367-009-0101-9"),
    ("jiang2022", "G. Jiang, W. Yu, H. Lei, Novel solar membrane distillation system based on Ti{{s|3}}C{{s|2}}T{{s|X}} MXene nanofluids "
     "with high photothermal conversion efficiency, Desalination 539 (2022) 115930. https://doi.org/10.1016/j.desal.2022.115930"),
    ("kim2012", "E. Kim, K. Osseo-Asare, Dissolution windows for hydrometallurgical purification of metallurgical-grade silicon to "
     "solar-grade silicon: Eh-pH diagrams for Fe silicides, Hydrometallurgy 127-128 (2012) 178-186. "
     "https://doi.org/10.1016/j.hydromet.2012.05.013"),
    ("qian2023", "X. Qian, Z. Wang, J. Sun, F. Zeng, D. Wu, Z. Xie, et al., Tuning the oxidation level of GO to construct high "
     "performance crosslinked lamellar graphene oxide membrane for pervaporation desalination, Desalination 568 (2023) 117015. "
     "https://doi.org/10.1016/j.desal.2023.117015"),
    ("ren2015", "C.E. Ren, K.B. Hatzell, M. Alhabeb, Z. Ling, K.A. Mahmoud, Y. Gogotsi, Charge- and size-selective ion sieving through "
     "Ti{{s|3}}C{{s|2}}T{{s|x}} MXene membranes, The Journal of Physical Chemistry Letters 6 (2015) 4026-4031. "
     "https://doi.org/10.1021/acs.jpclett.5b01895"),
    ("tan2021", "Y. Tan, T. Jiang, G.Z. Chen, Mechanisms and product options of magnesiothermic reduction of silica to silicon for "
     "lithium-ion battery applications, Frontiers in Energy Research 9 (2021) 651386. https://doi.org/10.3389/fenrg.2021.651386"),
    ("tian2022", "Y. Tian, L. Wang, B. Yang, Y. Dai, B. Xu, F. Wang, et al., Comparative evaluation of energy and resource consumption "
     "for vacuum carbothermal reduction and Pidgeon process used in magnesium production, Journal of Magnesium and Alloys 10 "
     "(2022) 697-706. https://doi.org/10.1016/j.jma.2020.09.024"),
    ("zhao2026", "X. Zhao, J. Tan, B. Fan, R.A. Soomro, H. Cui, N. Qiao, et al., Long-term aqueous stability of Ti{{s|3}}C{{s|2}}T{{s|x}} "
     "MXene achieved via a synergistic corrosion inhibition strategy, Corrosion Science 265 (2026) 113823. "
     "https://doi.org/10.1016/j.corsci.2026.113823"),
]
N0 = 26                                        # the manuscript ends at [26]
NUM = {k: N0 + i + 1 for i, (k, _) in enumerate(NEW_REFS)}


def R(*keys):
    return "[" + ",".join(str(NUM[k]) for k in keys) + "]"


# ------------------------------------------------------------------ the new text
INTRO_SENTENCE = (
    " The discussion also identifies downstream uses for the recovered minerals, including a preliminary research direction "
    "involving magnesium-mediated silicon production, and keeps these unmodeled extensions outside the quantitative system "
    "boundary."
)

S53 = [
    "A fourth direction, at an earlier stage of development, is the use of two-dimensional laminate membranes in place of "
    "conventional crystallization. Crosslinked graphene oxide (GO) and MXene laminates are of interest because their interlayer "
    "structure can be modified to influence water transport, ion exclusion and material stability, and because they might be "
    "integrated with membrane distillation or pervaporation under the hypersaline conditions of this architecture. Confining the "
    "interlayer spacing of GO laminates physically suppresses the swelling that otherwise limits them in water and has given 97% "
    "NaCl rejection " + R("abraham2017") + ", and borate crosslinking of highly oxidized GO has been used to improve the stability "
    "and permselectivity of lamellar GO membranes for pervaporation desalination " + R("qian2023") + ". MXene laminates sieve ions "
    "by hydration radius and charge at high water flux " + R("ren2015") + " and have been applied to the production of drinking "
    "water from seawater " + R("ding2020") + "; their surface chemistry is tunable and their photothermal response has been used to "
    "drive solar membrane distillation " + R("jiang2022") + ", which could in principle supply part of the thermal duty from "
    "sunlight. Their stability in aqueous media over long periods is itself an active problem " + R("zhao2026") + ".",

    "Neither material can presently be assumed to outperform mechanical vapor compression under the conditions modeled here. The "
    "reported performance belongs to individual laboratory membranes and does not establish the energy consumption, durability or "
    "economic viability of an integrated industrial crystallization stage. A comparison against the MVC reference used here would "
    "have to be made at the modeled crystallizer-feed salinity of 124.8 g L{{u|−1}} (Table 1) and would have to report water flux, "
    "salt rejection, scaling resistance, structural stability over extended operation, membrane replacement requirements and "
    "performance in a representative multicomponent brine rather than in a single salt. A complete system comparison would then "
    "have to include thermal-energy input, auxiliary electricity, membrane area, capital expenditure and the management of the "
    "residual concentrated brine and crystallized salts; a solar-assisted configuration would additionally have to account for "
    "solar intermittency, collection area and any supplementary heating or vacuum duty. Those measurements would establish whether "
    "advanced laminate membranes offer a genuine system-level advantage rather than improved laboratory-scale membrane "
    "performance. Until they exist, no rejection figure, photothermal conversion efficiency or elimination of crystallizer "
    "electricity is claimed or assumed here.",
]

S55_HEAD = "5.5 Future research: brine-derived magnesium for metallothermic silicon production"
S55 = [
    "A second research direction is whether magnesium recovered from desalination brine could serve as a precursor for "
    "metallothermic silicon production. Unlike the mineral valorization modeled above, which treats recovered magnesium hydroxide "
    "as a saleable product, this pathway would evaluate its downstream conversion into a higher-value material. It is proposed here "
    "as a separate research hypothesis rather than as an improvement demonstrated by the present model, and none of it enters the "
    "mass, energy, carbon or economic balances of Sections 3 and 4.",

    "The conceptual pathway recovers magnesium hydroxide from the concentrated brine, converts it into a suitable "
    "magnesium-containing intermediate, produces elemental magnesium and then uses that metal to reduce externally supplied "
    "silica. The principal reduction reaction is SiO{{s|2}} + 2 Mg → Si + 2 MgO, which has been studied in detail as a route to "
    "porous silicon " + R("tan2021") + ".",

    "The reaction is chemically established, but integrating it into a desalination facility introduces substantial additional "
    "processing requirements. Magnesium recovered as the hydroxide cannot substitute directly for elemental magnesium: conversion "
    "to the metal requires a separate reduction process, such as molten-salt electrolysis or thermal reduction, and the dominant "
    "industrial route is itself energy- and carbon-intensive, reported at 8.68 t of coal equivalent and 26.3 t CO{{s|2}}-eq per tonne "
    "of magnesium against 2.68 t of coal equivalent for a vacuum carbothermal alternative that is not yet industrially practiced "
    + R("gao2009", "tian2022") + ". The resulting silicon would also require separation, purification and product-quality "
    "verification, and the purification of metallurgical-grade silicon to solar grade is a process problem in its own right "
    + R("kim2012") + ".",

    "The principal research question is whether the additional silicon value could compensate for the energy, capital, reagent and "
    "purification requirements of the complete conversion pathway. Any assessment must retain the upstream magnesium-recovery cost "
    "established in Section 4.7 and must account for the mineral revenue forgone when recovered magnesium is consumed internally "
    "rather than sold, which the model carries as the mineral credit of Table 8. A future investigation would proceed in four "
    "stages: close a magnesium and silicon mass balance using the recovered quantities the present model predicts; determine the "
    "minimum and practical energy requirements for magnesium-metal production, silica reduction and downstream purification; "
    "extend the technoeconomic and carbon-accounting boundaries to the additional equipment, consumables, waste streams and "
    "product-quality requirements; and compare direct mineral sales with silicon production through sensitivity analyses under "
    "consistent electricity-price, carbon-intensity and capital-cost assumptions.",

    "Product quality is a major uncertainty. Producing elemental silicon does not establish that solar-grade or "
    "semiconductor-grade specifications can be achieved, and a future economic case should distinguish demonstrated purity from "
    "hypothetical high-purity scenarios. This pathway would not eliminate the alkalinity requirement identified here, nor would it "
    "resolve the calcium balance without further process changes; its contribution would be to test whether a downstream use for "
    "recovered magnesium could improve the economics of an otherwise chemically consistent brine-valorization architecture. No "
    "silicon revenue, energy saving or carbon credit is included in the results reported in this paper.",
]

S56_ADD = [
    "Advanced GO- and MXene-based membrane configurations are identified in Section 5.3 as potential alternatives for future "
    "investigation, but their performance under the modeled hypersaline operating conditions has not been experimentally "
    "established and is not incorporated into the system balances reported here.",

    "The proposed magnesium-to-silicon pathway of Section 5.5 is not represented in the mass, energy, carbon or economic balances "
    "of this study. Its feasibility depends on additional magnesium-metal production and silicon purification processes, neither of "
    "which is evaluated here. The pathway should therefore be interpreted exclusively as a direction for subsequent research, not "
    "as evidence that the economic or environmental constraints identified in this study can be overcome.",
]

S57_ADD = [
    "An additional experimental program could evaluate crosslinked GO and MXene-based membranes for hypersaline separation. "
    "Candidate materials should first be screened against representative crystallizer-feed compositions and then operated under "
    "realistic temperature and concentration conditions for an extended period, reporting water flux, salt rejection, scaling and "
    "fouling behavior, membrane degradation, specific thermal-energy demand and auxiliary electricity consumption. Only "
    "configurations demonstrating adequate stability and separation performance would justify a system-level energy, carbon and "
    "economic comparison against the MVC reference.",

    "The silicon pathway belongs to a separate subsequent study: close its material balance first, then bound the energy of "
    "magnesium-metal production and silicon purification, and only then evaluate life-cycle carbon and economics. None of those "
    "tests was performed here.",
]

CONCL_ADD = (
    "Two further directions follow from this analysis without altering any result reported in it. Whether advanced "
    "two-dimensional membrane materials, including crosslinked graphene oxide and MXene laminates, can reduce the energy or "
    "environmental burden of final brine concentration and crystallization without introducing unacceptable stability, scaling or "
    "capital-cost penalties remains to be established. The conversion of recovered magnesium into a reductant for silicon "
    "production is a separate downstream opportunity whose relevance to desalination economics is undetermined, and establishing "
    "it would require a complete accounting of magnesium-metal production, silicon purification, displaced mineral revenue and the "
    "associated environmental burdens."
)

CLOSING_SENTENCE = ("This paper does not report a working valorization process; it identifies the conditions such a process would "
                    "have to meet and quantifies how far the architecture examined here is from them.")

# ------------------------------------------------------------------ helpers
TOK = re.compile(r"(\{\{[su]\|.*?\}\})")


def add_rich(par, text, template_run=None):
    """Append text to a paragraph, honouring {{s|x}} subscript and {{u|x}} superscript markup and
    copying the character formatting of a template run so the insertion matches the surrounding text."""
    for piece in TOK.split(text):
        if not piece:
            continue
        kind = None
        if piece.startswith("{{s|"):
            kind, piece = "sub", piece[4:-2]
        elif piece.startswith("{{u|"):
            kind, piece = "sup", piece[4:-2]
        r = par.add_run(piece)
        if template_run is not None:
            r.font.size = template_run.font.size
            r.font.name = template_run.font.name
            if template_run.font.name:
                r._element.rPr.rFonts.set(qn("w:eastAsia"), template_run.font.name)
            r.bold = template_run.bold
            r.italic = template_run.italic
        if kind == "sub":
            r.font.subscript = True
        elif kind == "sup":
            r.font.superscript = True
    return par


def clone_after(anchor_par, doc, fmt_par=None):
    """A new empty paragraph immediately after anchor_par, taking its paragraph formatting from
    fmt_par (default: anchor_par), so a heading inserted after body text is still formatted as a heading."""
    new_p = copy.deepcopy((fmt_par or anchor_par)._p)
    for child in list(new_p):
        if child.tag != qn("w:pPr"):
            new_p.remove(child)
    anchor_par._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    return Paragraph(new_p, anchor_par._parent)


def insert_block(anchor, texts, doc, template_run, fmt_par=None):
    """Insert paragraphs after `anchor`, in order; returns the last one inserted. fmt_par supplies the
    paragraph formatting for the first insertion when the anchor is a heading rather than body text."""
    cur = anchor
    for i, t in enumerate(texts):
        cur = clone_after(cur, doc, fmt_par=fmt_par if i == 0 else None)
        add_rich(cur, t, template_run)
    return cur


def find(doc, pred, what):
    for i, p in enumerate(doc.paragraphs):
        if pred(p):
            return i, p
    raise SystemExit("not found: " + what)


def body_run(p):
    for r in p.runs:
        if r.text.strip():
            return r
    return None


# ------------------------------------------------------------------ apply
doc = Document(SRC)

# 1. Introduction: one sentence at the end of the last Introduction paragraph
i, p = find(doc, lambda p: p.text.startswith("This study addresses that gap for one representative architecture"), "intro para")
add_rich(p, INTRO_SENTENCE, body_run(p))
print("intro sentence appended to P%03d" % i)

# 2. Section 5.3: two paragraphs after the existing one
i, p = find(doc, lambda p: p.text.startswith("Mechanical vapor compression dominates the energy account"), "5.3 body")
tmpl = body_run(p)
last53 = insert_block(p, S53, doc, tmpl)
print("5.3 extended after P%03d" % i)

# 3. New Section 5.5, after the last paragraph of 5.4
i, p = find(doc, lambda p: p.text.startswith("Five directions follow from the analysis"), "5.4 last para")
head_tmpl_i, head_tmpl = find(doc, lambda p: p.text.strip() == "5.4 Toward a closed valorization loop", "5.4 heading")
h = clone_after(p, doc, fmt_par=head_tmpl)
hr = h.add_run(S55_HEAD)
ht = head_tmpl.runs[0]
hr.bold, hr.font.size, hr.font.name = True, ht.font.size, ht.font.name
insert_block(h, S55, doc, tmpl, fmt_par=p)
print("new 5.5 inserted after P%03d" % i)

# 4. Renumber 5.5 -> 5.6 and 5.6 -> 5.7, and extend each
i, p = find(doc, lambda p: p.text.strip() == "5.5 Limitations", "old 5.5 heading")
p.runs[0].text = "5.6 Limitations"
i, p = find(doc, lambda p: p.text.startswith("The osmotic coefficient correction in Eq. (5)"), "5.6 body")
insert_block(p, S56_ADD, doc, body_run(p))
print("5.6 Limitations renumbered and extended")

i, p = find(doc, lambda p: p.text.strip() == "5.6 Proposed validation", "old 5.6 heading")
p.runs[0].text = "5.7 Proposed validation"
i, p = find(doc, lambda p: p.text.startswith("Three measurements would settle the questions"), "5.7 body")
insert_block(p, S57_ADD, doc, body_run(p))
print("5.7 Proposed validation renumbered and extended")

# 5. Conclusions: split the closing paragraph and insert the forward-looking paragraph before its last sentence
i, p = find(doc, lambda p: p.text.startswith("Treated as one design problem"), "conclusions last para")
assert p.text.rstrip().endswith(CLOSING_SENTENCE), "closing sentence not where expected"
tmpl_c = body_run(p)
# strip the closing sentence off the existing paragraph, run by run
remaining = CLOSING_SENTENCE
for r in reversed(p.runs):
    if not remaining:
        break
    if r.text.endswith(remaining):
        r.text = r.text[: -len(remaining)].rstrip()
        remaining = ""
    elif remaining.endswith(r.text):
        remaining = remaining[: -len(r.text)].rstrip()
        r.text = ""
assert not remaining, "could not split the closing sentence"
new_par = insert_block(p, [CONCL_ADD, CLOSING_SENTENCE], doc, tmpl_c)
print("conclusions paragraph inserted before the closing sentence")

# 6. Cross-reference: Sections 3 and 5.5 -> 5.6
n_fix = 0
for p in doc.paragraphs:
    for r in p.runs:
        if "Sections 3 and 5.5" in r.text:
            r.text = r.text.replace("Sections 3 and 5.5", "Sections 3 and 5.6")
            n_fix += 1
print("cross-references updated:", n_fix)
assert n_fix == 1, "expected exactly one 5.5 cross-reference"

# 7. References
i, p = find(doc, lambda p: p.text.startswith("[26] F.J. Millero"), "last reference")
ref_tmpl = body_run(p)
cur = p
for key, entry in NEW_REFS:
    cur = clone_after(cur, doc)
    add_rich(cur, "[%d] %s" % (NUM[key], entry), ref_tmpl)
print("references added:", len(NEW_REFS), "->", N0 + len(NEW_REFS))

doc.save(OUT)

# ------------------------------------------------------------------ checks
chk = Document(OUT)
text = "\n".join(p.text for p in chk.paragraphs)
for t in chk.tables:
    for row in t.rows:
        for c in row.cells:
            text += "\n" + c.text
assert "{{" not in text and "}}" not in text, "markup leaked into the document"
for n in range(1, N0 + len(NEW_REFS) + 1):
    assert "[%d] " % n in text, "reference %d missing from the list" % n
cited = set()
for m in re.finditer(r"\[(\d+(?:,\d+)*)\]", text):
    for x in m.group(1).split(","):
        cited.add(int(x))
missing = [NUM[k] for k, _ in NEW_REFS if NUM[k] not in cited]
assert not missing, "new references never cited: %s" % missing
heads = [p.text.strip() for p in chk.paragraphs if re.match(r"^5\.\d ", p.text.strip())]
print("section headings:", heads)
assert heads == ["5.1 Comparison with conventional SWRO", "5.2 Comparison with the brine-valorization literature",
                 "5.3 Alternative crystallization pathways", "5.4 Toward a closed valorization loop", S55_HEAD,
                 "5.6 Limitations", "5.7 Proposed validation"], "section numbering is wrong"
print("saved", OUT)
print("paragraphs %d -> %d, tables %d, figures unchanged"
      % (len(Document(SRC).paragraphs), len(chk.paragraphs), len(chk.tables)))
