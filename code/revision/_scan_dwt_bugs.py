# -*- coding: utf-8 -*-
"""Scan the DWT documents for the bug classes found in the Wazoku white paper.

Three classes, all of which were real there:

  1  double-escaped unicode, which prints as a literal backslash-u sequence;
  2  doubled percent signs in strings that are not %-formatted;
  3  a quantity stated on the wrong basis -- per cubic metre of FEED where the
     sentence claims per cubic metre of PERMEATE. That one was a factor of 2.4
     and is the only one a reader would notice.

Also flags leftover placeholders and hardcoded numerics in the builders that
should be read from the model JSON.

    python _scan_dwt_bugs.py
"""
import io
import os
import re

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
DWT = r"C:\Users\Leon\Downloads\DWT"
BS = chr(92)

DOCS = [os.path.join(DWT, "ILEDBV_Manuscript_Revised_v3.docx"),
        os.path.join(DWT, "DWT_Response_to_Reviewers.docx")]
BUILDERS = ["build_full_manuscript.py", "build_response_letter.py",
            "build_wazoku_routeD.py"]

report = []


def text_of(path):
    d = Document(path)
    return "\n".join([p.text for p in d.paragraphs] +
                     [c.text for t in d.tables for row in t.rows for c in row.cells])


# ------------------------------------------------- 1 and 2: rendering artefacts
print("=" * 72)
print("1-2. RENDERING ARTEFACTS IN THE BUILT DOCUMENTS")
print("=" * 72)
for path in DOCS:
    if not os.path.exists(path):
        print("   MISSING %s" % os.path.basename(path))
        continue
    t = text_of(path)
    lit = re.findall(re.escape(BS) + r"u[0-9a-fA-F]{4}", t)
    pct = t.count("%%")
    brackets = re.findall(r"\[[A-Za-z][^\]]{2,60}\]", t)
    placeholders = [b for b in brackets
                    if not re.match(r"\[\d", b) and "," in b or "insert" in b.lower()]
    print("   %-40s backslash-u %d, %%%% %d, placeholders %d"
          % (os.path.basename(path)[:40], len(lit), pct, len(placeholders)))
    for b in placeholders[:5]:
        report.append("placeholder in %s: %s" % (os.path.basename(path), b))
    for m in re.finditer(r".{40}%%.{40}", t):
        report.append("doubled percent in %s: ...%s..."
                      % (os.path.basename(path), m.group(0).replace("\n", " ")))

# --------------------------------------------- 3: feed vs permeate basis errors
print("\n" + "=" * 72)
print("3. BASIS ERRORS (per m3 FEED stated as per m3 PERMEATE)")
print("=" * 72)
# the model's own constants, and what they mean
KG_SODA, KG_LIME = 1.0231, 3.7923          # per m3 FEED
R = 0.45
feed_vals = {"soda ash": KG_SODA, "lime": KG_LIME,
             "soda ash + lime": KG_SODA + KG_LIME}
print("   reference: soda %.4f, lime %.4f kg per m3 FEED" % (KG_SODA, KG_LIME))
print("              equivalently %.3f, %.3f kg per m3 PERMEATE"
      % (KG_SODA / R, KG_LIME / R))
print("              combined %.3f feed, %.3f permeate"
      % (KG_SODA + KG_LIME, (KG_SODA + KG_LIME) / R))

SUSPECT = {}
for label, v in feed_vals.items():
    SUSPECT["%.1f" % v] = (label, "feed-basis value")
    SUSPECT["%.2f" % v] = (label, "feed-basis value")

for path in DOCS:
    if not os.path.exists(path):
        continue
    t = text_of(path)
    name = os.path.basename(path)
    for m in re.finditer(r"([\d.]+)\s*kg[^.]{0,80}?(permeate|feed)", t, re.I):
        val, basis = m.group(1), m.group(2).lower()
        if val in SUSPECT and basis == "permeate":
            report.append("BASIS: %s states %s kg per m3 permeate, but %s is a "
                          "%s" % (name, val, val, SUSPECT[val][1]))
            print("   !! %s: '%s kg ... permeate' looks like a feed-basis number"
                  % (name, val))
print("   (no basis mismatches found)" if not any(
    r.startswith("BASIS") for r in report) else "")

# ------------------------------------- hardcoded numerics in the builder sources
print("\n" + "=" * 72)
print("4. HARDCODED NUMERICS IN BUILDER SOURCES")
print("=" * 72)
for b in BUILDERS:
    p = os.path.join(HERE, b)
    if not os.path.exists(p):
        continue
    src = io.open(p, encoding="utf-8").read()
    # numbers written straight into a display string rather than %-formatted
    hits = []
    for m in re.finditer(r'"[^"]*?(\d+\.\d+)\s*(kWh|kg|\$|m-3|%)[^"]*?"', src):
        frag = m.group(0)
        if "%." in frag:
            continue
        hits.append(m.group(1) + " " + m.group(2))
    uniq = sorted(set(hits))
    print("   %-32s %d literal value(s)%s"
          % (b[:32], len(uniq), (": " + ", ".join(uniq[:8])) if uniq else ""))
    for u in uniq:
        report.append("hardcoded in %s: %s" % (b, u))

io.open(os.path.join(HERE, "_dwt_bug_report.txt"), "w", encoding="utf-8").write(
    "\n".join(report) if report else "clean")
print("\n" + "=" * 72)
print("%d item(s) written to _dwt_bug_report.txt" % len(report))
