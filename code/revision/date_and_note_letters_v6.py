# -*- coding: utf-8 -*-
"""Date the correspondence consistently with the manuscript it describes, and record the final pass.

The cover letter and the response letter were written on 2 October, before the two future-research
subsections and the final corrections of 3 October. Both are redated and both gain one entry covering
the corrections, so the letters describe the version actually submitted.

    python date_and_note_letters_v6.py
"""
import copy
import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.abspath(os.path.join(HERE, "..", "..", "revision2"))

NOTE = (
    "A final editorial pass makes four corrections, none of which changes a result. Section 5.3 now describes the laminate "
    "membranes as potential components of alternative brine-concentration and crystallization systems, and states that membrane "
    "separation and complete crystallization are not interchangeable. The electricity-price crossings of Section 4.12 and the "
    "Conclusions are labelled modeled thresholds under the stated cost assumptions rather than demonstrated commercial break-even "
    "prices, since the capital cost of the electrochemical unit and its coproduct handling lie outside the boundary. The "
    "generative-AI declaration is rewritten to state what the tool did and what was verified, and by what. Data availability now "
    "names the plotting script for Figures 2 and 7 and the release that contains it, because that script was added after v2.3.0, "
    "which remains the tag cited for the model code."
)


def redate(doc):
    n = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if "2 October 2026" in r.text:
                r.text = r.text.replace("2 October 2026", "3 October 2026")
                n += 1
    return n


def bullet_after(doc, start, text):
    i = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith(start))
    src = doc.paragraphs[i]
    new_p = copy.deepcopy(src._p)
    src._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    par = Paragraph(new_p, src._parent)
    for r in list(par.runs)[1:]:
        r._r.getparent().remove(r._r)
    par.runs[0].text = text
    par.runs[0].bold = False


for name, anchor in (("DWT_Revision2_Cover_Letter.docx", "Two forward-looking subsections"),
                     ("DWT_Response_to_Reviewers_R2.docx", "Two forward-looking subsections")):
    f = os.path.join(REV, name)
    doc = Document(f)
    n = redate(doc)
    bullet_after(doc, anchor, NOTE)
    doc.save(f)
    print("%s: %d dates updated, final-pass entry added" % (name, n))

for name in ("DWT_Revision2_Cover_Letter.docx", "DWT_Response_to_Reviewers_R2.docx"):
    t = "\n".join(p.text for p in Document(os.path.join(REV, name)).paragraphs)
    print(name, "| '2 October 2026' left:", t.count("2 October 2026"), "| dated 3 October:", t.count("3 October 2026"),
          "| names v6:", t.count("_v6."))
