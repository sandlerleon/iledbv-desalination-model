# -*- coding: utf-8 -*-
"""The rewritten Introduction and reference list for DWT-D-26-01366.

Reviewers 1, 3 and 4 all asked for the same thing: continuous prose rather than
subsections, a critical review of recent progress in each of the five component
technologies rather than definitions of them, an account of which pairs have
already been integrated and what that revealed, and a research gap stated
plainly. Citation numbers refer to _refs_final.json, every entry of which was
verified against Crossref.

    python build_introduction.py
"""
import io, json, os
from docx import Document
from docx.shared import Pt, Inches

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\DWT\DWT_Revised_Introduction_v1.docx"
REFS = json.load(io.open(os.path.join(HERE, "_refs_final.json"), encoding="utf-8"))
V2 = json.load(io.open(os.path.join(HERE, "iledbv_revision_v2.json"), encoding="utf-8"))

w_min = V2["t1_concentrator"]["least_work_kwh_per_m3_brine"]
A = V2["precipitation_routes"]["A: soda ash + lime (as written)"]

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.9)


def h(t, size=13, before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(size)


def p(t, italic=False, size=11, after=10, first_indent=True):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(after)
    if first_indent:
        par.paragraph_format.first_line_indent = Inches(0.25)
    r = par.add_run(t); r.italic = italic; r.font.size = Pt(size)


h("Revised Introduction and references — DWT-D-26-01366", 15, 0)
p("Continuous prose, no subsections (R1-1). Every citation verified against "
  "Crossref; the numbering matches the list at the end.", italic=True, size=9.5,
  first_indent=False)

h("1.  Introduction")

p("Seawater reverse osmosis (SWRO) supplies a growing share of municipal water in "
  "arid coastal regions, and its energy performance has improved to the point where "
  "further gains at the membrane stage are incremental. Specific energy consumption "
  "(SEC) in well-designed modern plants equipped with isobaric energy recovery falls "
  "in the range 2.5-3.5 kWh m-3, and a recent audited measurement on an operating "
  "2,500 m3 d-1 plant reports 1.794 kWh m-3 under a low-flux design with new "
  "membranes [16]; a survey across operating plants places the practical "
  "distribution firmly in that band [24]. The thermodynamic floor for seawater at "
  "typical recovery is lower still [18]. These figures matter for how any new "
  "architecture should be assessed: the headroom between a modern plant and the "
  "thermodynamic limit is now small enough that claims of large energy savings at "
  "the desalination stage should be treated sceptically, and a proposed architecture "
  "is better judged on what it does with the concentrate than on what it saves at "
  "the membrane.", first_indent=False)

p("That concentrate is the unresolved problem. Every cubic metre of permeate is "
  "accompanied by roughly an equal volume of brine at approximately twice the feed "
  "salinity, and its disposal is increasingly constrained. Two responses have "
  "developed in parallel. The first treats the brine as a waste to be minimised, "
  "through minimal-liquid-discharge (MLD) and zero-liquid-discharge (ZLD) "
  "configurations; module-scale modelling of enhanced-RO pathways shows how far "
  "membrane processes can carry the concentration duty before a thermal step becomes "
  "unavoidable [2], and comparative techno-economic assessment of MLD against full "
  "ZLD shows the cost penalty rising steeply over the final increments of water "
  "recovery [15]. The second treats the brine as an ore, recovering calcium, "
  "magnesium and in some proposals lithium as saleable co-products. The two "
  "responses are complementary in principle and are seldom combined in practice.")

p("Within the desalination stage itself, two technologies are mature enough that "
  "their behaviour is well characterised. Isobaric pressure-exchanger energy "
  "recovery is standard practice, with device efficiencies above 95% and a design "
  "literature going back two decades [11,23]; it is the single change most "
  "responsible for the SEC figures quoted above, and it leaves little further to "
  "recover. Subsurface intake, through beach wells or seabed galleries, is less "
  "universally applicable but well studied: comparative assessments against open "
  "intakes report lower pretreatment demand, reduced chemical dosing and lower "
  "entrainment impact [21], operating-cost comparisons reach similar conclusions "
  "where the hydrogeology permits [7], and numerical modelling has clarified the "
  "conditions under which a beach well can sustain design yield without clogging "
  "[1]. The limitation is site dependence rather than process performance, and it is "
  "the reason subsurface intake is a design choice rather than a general "
  "recommendation.")

p("The brine-concentration step is where the literature is least settled, and where "
  "the present study finds the largest discrepancy. Osmotically assisted and "
  "high-pressure membrane configurations extend membrane concentration well beyond "
  "conventional SWRO brine salinity [3], but the energy demand rises sharply because "
  "the osmotic pressure to be overcome rises with it. Panagopoulos compared minimum "
  "and actual energy consumption across brine-treatment technologies and found the "
  "gap between the two to be both large and frequently unacknowledged [14]. That "
  "comparison is the appropriate reference point for any assumed concentrator "
  "energy, and, as Section 4.5 shows, values in common use in integrated-architecture "
  "studies - including the value adopted in the first version of this work - fall "
  "below the least work of separation for the feed concerned (%.2f kWh m-3), which "
  "no process can do." % w_min)

p("For the final concentration and crystallization duty, mechanical vapour "
  "compression (MVC) remains the reference technology, and techno-economic analysis "
  "of MVC in near-ZLD service establishes both its energy intensity and the "
  "conditions under which it dominates plant operating cost [20]. Three alternatives "
  "are under active development and deserve explicit comparison rather than "
  "omission. Membrane distillation crystallization has been demonstrated on "
  "hypersaline feeds with seeding to control scaling [8,19]; eutectic freeze "
  "crystallization offers a route to separate salts and water at lower theoretical "
  "energy, and has been applied specifically to RO brine [13]; and solar-driven "
  "interfacial evaporation has advanced rapidly, with architectures now able to "
  "sustain evaporation from high-salinity brine while harvesting the crystallized "
  "salt autonomously rather than fouling [4,5,6]. The last of these substitutes land "
  "area and capital for electricity, which changes the carbon comparison materially "
  "and is taken up in Section 4.7.")

p("Mineral recovery from the concentrated brine is the component with the widest "
  "spread between laboratory result and process claim. Sequential precipitation of "
  "calcium and magnesium is well established chemically, and recent work has "
  "quantified how the sodium chloride background affects yield, purity, particle "
  "morphology and carbonate polymorph - all of which determine whether a recovered "
  "solid meets any industrial specification [9]. Product-quality-directed studies "
  "have produced high-purity magnesium sulfate and vaterite-type calcium carbonate "
  "from seawater and brine feeds [17], establishing that specification-grade product "
  "is attainable but not that it is attainable as a by-product of a plant optimised "
  "for water. Mineral carbonation has been proposed as the route to selective "
  "recovery [10], and magnesium recovery via CO2 mineralization has been demonstrated "
  "with carbon dioxide supplying the carbonate [25]. This last result is directly "
  "relevant to the present architecture: it indicates that the carbonate source and "
  "the alkalinity source are design variables rather than fixed reagents, which is "
  "the observation Section 4.7 develops.")

p("Integration of these components has been attempted only in part. Energy recovery "
  "and subsurface intake are routinely combined, because neither constrains the "
  "other. Brine concentration has been coupled to crystallization in MLD and ZLD "
  "studies [2,15]. Mineral recovery has been studied on brine streams in isolation "
  "from the plant that produced them [9,10,17]. What has not been reported is an "
  "end-to-end assessment in which the intake, the energy-recovery configuration, the "
  "concentration duty, the precipitation chemistry and the crystallization load are "
  "solved on a single mass, salt, energy and carbon balance - and, critically, in "
  "which each stage is then tested against the physical bound it must respect. The "
  "absence matters because the components interact: the concentrator sets the "
  "crystallizer feed salinity, the precipitation chemistry returns ions to the stream "
  "the concentrator just removed, and the reagents that drive the precipitation carry "
  "an embodied carbon burden that no stage-level analysis sees.")

p("This study addresses that gap for one representative architecture, referred to "
  "here as ILEDBV (Integrated Low-Energy Desalination with Brine Valorization), "
  "combining subsurface intake, isobaric pressure-exchanger energy recovery, "
  "membrane brine concentration, selective Ca/Mg precipitation and MVC "
  "crystallization. Figure 1 is the author's conceptual integration of these five "
  "published component technologies rather than a description of an existing "
  "installation, and is presented as such. The objective is not to advocate the "
  "architecture but to test it: each subsystem claim is checked against a "
  "thermodynamic or stoichiometric constraint rather than carried as an assumed "
  "value, and a claim is allowed to fail. Three questions are posed. What is the "
  "modelled specific energy requirement once each stage is held to its physical "
  "bound, and how does it decompose? What governs the recovery of minerals from the "
  "brine, given that the reagents required to precipitate them are themselves "
  "products of energy-intensive processes? And under what combination of "
  "capital cost, electricity price and realisable mineral value, propagated "
  "probabilistically rather than varied one factor at a time, does the integrated "
  "architecture outperform conventional SWRO? The answers are, respectively, that "
  "the energy advantage does not survive a modern comparator, that mineral recovery "
  "is limited by alkalinity rather than by mineral inventory at a cost of %.1f kg "
  "CO2 per cubic metre of permeate, and that the economic case is close to evenly "
  "balanced." % A["co2_reagents"])

# ------------------------------------------------------------------ references
h("References (verified)")
p("Each entry was retrieved from Crossref and checked for title, authors, journal, "
  "volume and pages. Supplementary-material records and preprint shadows returned by "
  "bibliographic search were excluded.", italic=True, size=9.5, first_indent=False)

for i, r in enumerate(REFS, 1):
    a = ", ".join(r["authors"][:6]) + (", et al." if len(r["authors"]) > 6 else "")
    vol = (" %s" % r["volume"]) if r["volume"] else ""
    pg = (" (%s) %s" % (r["year"], r["pages"])) if r["pages"] else " (%s)" % r["year"]
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    run = par.add_run("[%d] %s, %s, %s%s%s. https://doi.org/%s"
                      % (i, a, r["title"], r["journal"], vol, pg, r["doi"]))
    run.font.size = Pt(9.5)

# also emit the paragraphs so the full-manuscript builder can reuse them verbatim
intro = [x.text for x in doc.paragraphs
         if len(x.text.split()) > 40 and not x.text.startswith("[")]
io.open(os.path.join(HERE, "_intro_paras.json"), "w", encoding="utf-8").write(
    json.dumps(intro, ensure_ascii=False, indent=1))
print("intro paragraphs emitted: %d" % len(intro))

doc.save(OUT)
words = sum(len(x.text.split()) for x in doc.paragraphs)
print("references: %d" % len(REFS))
print("introduction + references, words: %d" % words)
print("saved", OUT)
