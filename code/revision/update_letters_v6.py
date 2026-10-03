# -*- coding: utf-8 -*-
"""Bring the Revision 2 cover letter and response letter into line with manuscript v6.

Three things change. Section cross-references are renumbered where the new Section 5.5 displaced the
old ones (5.5 -> 5.6, 5.6 -> 5.7), and only inside an explicit "Section(s) ..." phrase, so that
numbers such as the 4.6-5.5 mg/L boron range are left alone. The enclosed file names move from v5 to
v6. And both letters gain a statement of what was added and that it changes no result, because the
two new subsections answer editorial guidance rather than a reviewer comment.

    python update_letters_v6.py
"""
import copy
import os
import re

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.abspath(os.path.join(HERE, "..", "..", "revision2"))
SEC = re.compile(r"Sections?\s+\d+(?:\.\d+)?(?:\s*(?:,|and|;)?\s*\d+(?:\.\d+)?)*")


def renumber(text):
    def fix(m):
        s = m.group(0)
        s = re.sub(r"(?<![\d.])5\.6(?![\d])", "5.<<7>>", s)
        s = re.sub(r"(?<![\d.])5\.5(?![\d])", "5.<<6>>", s)
        return s.replace("<<", "").replace(">>", "")
    return SEC.sub(fix, text)


def patch(doc, pairs):
    n = 0
    for p in doc.paragraphs:
        for r in p.runs:
            new = renumber(r.text)
            for a, b in pairs:
                new = new.replace(a, b)
            if new != r.text:
                r.text = new
                n += 1
    return n


def bullet_after(doc, idx, text):
    """Copy the list bullet at idx and put the new text in a fresh bullet beneath it."""
    src = doc.paragraphs[idx]
    new_p = copy.deepcopy(src._p)
    src._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    par = Paragraph(new_p, src._parent)
    for r in list(par.runs)[1:]:
        r._r.getparent().remove(r._r)
    par.runs[0].text = text
    par.runs[0].bold = False
    return par


def find(doc, start):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(start):
            return i
    raise SystemExit("not found: " + start)


COVER_BULLET = (
    "Two forward-looking subsections are added to the Discussion, neither of them requested by a reviewer. Section 5.3 now "
    "considers crosslinked graphene oxide and MXene laminate membranes as a possible alternative to the crystallization stage, "
    "and a new Section 5.5 sets out brine-derived magnesium as a precursor for metallothermic silicon production. Both are "
    "written as research directions and are excluded from the mass, energy, carbon and economic balances; the Limitations and "
    "Proposed validation sections, renumbered 5.6 and 5.7, record that exclusion. Ten references support them, [27] to [36], "
    "each resolved against its Crossref record."
)
RESP_BULLET = (
    "Two forward-looking subsections are added to the Discussion. These answer editorial guidance rather than a reviewer "
    "comment, and neither affects a result. Section 5.3 now considers crosslinked graphene oxide and MXene laminate membranes as "
    "a possible alternative to the crystallization stage, and a new Section 5.5 sets out brine-derived magnesium as a precursor "
    "for metallothermic silicon production. Both are stated as research directions; both are excluded from the mass, energy, "
    "carbon and economic balances, and the corresponding qualifications are recorded in the Limitations and in Proposed "
    "validation. Ten references supporting them are added as [27] to [36]. The former Sections 5.5 and 5.6 are accordingly "
    "renumbered 5.6 and 5.7."
)
NO_CHANGE = (" The two Discussion subsections added after that review are proposals for future work: they introduce no "
             "calculation and change no reported value.")

# ---------------------------------------------------------------- cover letter
f = os.path.join(REV, "DWT_Revision2_Cover_Letter.docx")
doc = Document(f)
n = patch(doc, [])
i = find(doc, "Sections 5.6 and 5.7 are extended")
bullet_after(doc, i, COVER_BULLET)
i = find(doc, "I should record plainly that no model")
doc.paragraphs[i].runs[-1].text = doc.paragraphs[i].runs[-1].text.rstrip() + NO_CHANGE
doc.save(f)
print("cover letter: %d runs renumbered, bullet and qualifier added" % n)

# ---------------------------------------------------------------- response letter
f = os.path.join(REV, "DWT_Response_to_Reviewers_R2.docx")
doc = Document(f)
n = patch(doc, [("ILEDBV_Manuscript_Revised_v5.docx", "ILEDBV_Manuscript_Revised_v6.docx"),
                ("ILEDBV_Manuscript_Revised_v5_marked.docx", "ILEDBV_Manuscript_Revised_v6_marked.docx")])
i = find(doc, "Sections 5.6 and 5.7 are extended")
bullet_after(doc, i, RESP_BULLET)
i = find(doc, "One point is worth stating at the outset")
doc.paragraphs[i].runs[-1].text = doc.paragraphs[i].runs[-1].text.rstrip() + NO_CHANGE
doc.save(f)
print("response letter: %d runs renumbered, bullet and qualifier added" % n)

# ---------------------------------------------------------------- check
for name in ("DWT_Revision2_Cover_Letter.docx", "DWT_Response_to_Reviewers_R2.docx"):
    d = Document(os.path.join(REV, name))
    txt = "\n".join(p.text for p in d.paragraphs)
    stale = [m.group(0) for m in SEC.finditer(txt) if re.search(r"(?<![\d.])5\.5(?![\d])", m.group(0)) and "renumbered" not in txt[max(0, m.start() - 90):m.start()]]
    print(name, "| v5 file names left:", txt.count("_v5.docx"), "| 'Section 5.5' spans left:", stale)
