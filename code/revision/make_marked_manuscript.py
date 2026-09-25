# -*- coding: utf-8 -*-
"""Marked-up copy of manuscript v4 for the DWT revision upload.

Every sentence of v4 that does not appear (near-verbatim) in the originally
submitted manuscript is highlighted yellow. Sentence-level highlighting is used
rather than tracked changes because the revision rewrites most of the paper, and
a tracked-change diff of a rewrite is unreadable.

    python make_marked_manuscript.py
"""
import copy
import difflib
import os
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.text.run import Run

HERE = os.path.dirname(os.path.abspath(__file__))
MS = os.path.abspath(os.path.join(HERE, "..", "..", "manuscript"))
ORIG = os.path.join(MS, "ILEDBV_Desalination_Manuscript.docx")
NEW = os.path.join(MS, "ILEDBV_Manuscript_Revised_v4.docx")
OUT = os.path.join(MS, "ILEDBV_Manuscript_Revised_v4_marked.docx")

SENT = re.compile(r"\S.*?(?:[.!?]+(?=\s|$)|$)")


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


orig_sents = set()
orig_text = []
for p in all_paragraphs(Document(ORIG)):
    for m in SENT.finditer(p.text):
        n = norm(m.group(0))
        if n:
            orig_sents.add(n)
            orig_text.append(n)


def is_old(s):
    n = norm(s)
    if not n or len(n) < 4 or n in orig_sents:
        return True
    close = difflib.get_close_matches(n, orig_text, n=1, cutoff=0.92)
    return bool(close)


def split_run(run, offset):
    """Split a text-only run at character offset; return the new second run."""
    text = run.text
    new_r = copy.deepcopy(run._r)
    run._r.addnext(new_r)
    run.text = text[:offset]
    second = Run(new_r, run._parent)
    second.text = text[offset:]
    return second


def text_only(run):
    return all(ch.tag in (qn("w:rPr"), qn("w:t")) for ch in run._r)


doc = Document(NEW)
new_chars = total_chars = 0
seen = set()
for p in all_paragraphs(doc):
    if id(p._p) in seen:
        continue
    seen.add(id(p._p))
    text = p.text
    ranges = [(m.start(), m.end()) for m in SENT.finditer(text) if not is_old(m.group(0))]
    total_chars += len(text.strip())
    if not ranges:
        continue
    new_chars += sum(b - a for a, b in ranges)
    # split text-only runs at range boundaries, then highlight runs inside ranges
    cuts = sorted({x for r in ranges for x in r})
    pos = 0
    for run in list(p.runs):
        L = len(run.text)
        if text_only(run):
            inner = [c - pos for c in cuts if pos < c < pos + L]
            cur, base = run, pos
            for c in inner:
                nxt = split_run(cur, c - (base - pos))
                base = pos + c
                cur = nxt
        pos += L
    pos = 0
    for run in p.runs:
        L = len(run.text)
        mid = pos + L / 2.0
        if L and any(a <= mid < b for a, b in ranges):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos += L
    assert p.text == text

pct = 100.0 * new_chars / total_chars
first = doc.paragraphs[0]
note = copy.deepcopy(first._p)
first._p.addprevious(note)
from docx.text.paragraph import Paragraph
note_p = Paragraph(note, first._parent)
for r in note_p.runs[1:]:
    r._r.getparent().remove(r._r)
note_p.runs[0].text = (
    "Marked-up copy for review: text that is new or changed relative to the original submission "
    "(DWT-D-26-01366) is highlighted in yellow. The revision rewrites most of the manuscript "
    "(%.0f%% of the text), so nearly all of it is highlighted; see the Response to Reviewers for "
    "a point-by-point account." % pct)
note_p.runs[0].font.highlight_color = None
note_p.runs[0].font.bold = False
note_p.runs[0].font.italic = True
note_p.runs[0].font.size = Pt(10)
doc.save(OUT)
print("wrote %s  (%.0f%% of text highlighted as new or changed)" % (OUT, 100.0 * new_chars / total_chars))
