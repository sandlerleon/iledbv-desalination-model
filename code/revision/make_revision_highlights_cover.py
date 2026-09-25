# -*- coding: utf-8 -*-
"""Highlights and revision cover letter for DWT-D-26-01366 (manuscript v4).

The Highlights and cover letter from the first submission quote results that the
revision withdrew (10.68 kWh m-3, 31% SEC advantage, $0.42 m-3 net LCOW), so they
cannot be carried forward. This writes replacements built from the model JSON, on
the original files as templates so the styling matches.

    python make_revision_highlights_cover.py
"""
import copy
import io
import json
import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
MS = os.path.abspath(os.path.join(HERE, "..", "..", "manuscript"))
V2 = json.load(io.open(os.path.join(HERE, "iledbv_revision_v2.json"), encoding="utf-8"))
AC = json.load(io.open(os.path.join(HERE, "alkalinity_calcium_results.json"), encoding="utf-8"))
mc = V2["monte_carlo"]
A = V2["precipitation_routes"]["A: soda ash + lime (as written)"]
ceil = AC["price_ceiling_usd_per_kwh"]["reversible_limit"]
closure = AC["optimisation"]["model base case"]["closure_cost_usd_per_m3"]

from docx import Document as _D
TITLE = _D(os.path.join(MS, "ILEDBV_Manuscript_Revised_v4.docx")).paragraphs[0].text.strip()
assert TITLE.startswith("Consistency Testing"), TITLE

HIGHLIGHTS = [
    "Assumed brine-concentrator energy lies below the least work of separation",
    "Alkalinity, not magnesium, limits Mg recovery from SWRO brine by precipitation",
    "Lime dosing releases %.1f times more calcium than the carbonate step removes"
    % A["ca_closure_ratio"],
    "Explicit reagent costing gives net LCOW of $%.2f/m3 vs $0.76/m3 for SWRO"
    % mc["lcow_net"]["median"],
    "No on-site base closes the calcium loop above $%.3f/kWh electricity" % ceil,
]
for h in HIGHLIGHTS:
    assert len(h) <= 85, (len(h), h)


def rewrite(src, dst, paras):
    doc = Document(src)
    body = doc.paragraphs
    templates = {}
    for p in body:
        templates.setdefault(p.style.name, p)
    for p in body[1:]:
        p._p.getparent().remove(p._p)
    anchor = body[0]
    first = True
    for style, text in paras:
        if first:
            tgt = anchor
            first = False
        else:
            el = copy.deepcopy(templates[style]._p)
            anchor._p.addnext(el)
            from docx.text.paragraph import Paragraph
            tgt = Paragraph(el, anchor._parent)
        runs = tgt.runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            tgt.add_run(text)
        anchor = tgt
    doc.save(dst)
    print("wrote", dst)


rewrite(os.path.join(MS, "Highlights_ILEDBV.docx"),
        os.path.join(MS, "Highlights_ILEDBV_Revised.docx"),
        [("Heading 1", "Highlights")] + [("List Bullet", h) for h in HIGHLIGHTS])

N = "Normal"
LETTER = [
    (N, "Leon Sandler"),
    (N, "Independent Researcher"),
    (N, "Northbrook, Illinois, USA"),
    (N, "sandler.leon@gmail.com"),
    (N, "September 25, 2026"),
    (N, ""),
    (N, "The Editor"),
    (N, "Desalination and Water Treatment"),
    (N, ""),
    (N, "Re: Revised submission of manuscript DWT-D-26-01366"),
    (N, ""),
    (N, "Dear Editor,"),
    (N, "Thank you for the opportunity to revise manuscript DWT-D-26-01366, originally titled "
        "\"Process and Technoeconomic Assessment of Integrated Subsurface-Intake SWRO, Isobaric "
        "Energy Recovery, and Near-Zero-Liquid-Discharge Brine Valorization.\" I am grateful to "
        "the four reviewers, whose comments exposed two errors that changed the paper's "
        "conclusions. The revision follows where those corrections led, and the title is now "
        "\"%s.\"" % TITLE),
    (N, "The principal changes are these. The brine-concentrator energy assumed in the original "
        "lay at or below the least work of separation for that feed; it is replaced by a bounded "
        "estimate, raising the modeled median total specific energy consumption to %.1f kWh m-3. "
        "The precipitation reagents are now costed explicitly: at $%.2f m-3 they exceed the "
        "realizable mineral credit, and Monte Carlo propagation gives a median net levelized cost "
        "of $%.2f m-3 against $0.76 m-3 for conventional SWRO. The claimed 31%% energy advantage "
        "is withdrawn, and the paper is reframed as a consistency test that identifies alkalinity "
        "supply and calcium closure as the binding constraints. A new Section 4.13 treats those "
        "two as a single linear optimization and quantifies the conditions under which the loop "
        "could close: closing the calcium balance costs $%.2f m-3 at the base case, and above "
        "$%.3f kWh-1 no on-site base can close it because the reversible limit of water "
        "dissociation binds. The paper does not report a working valorization process; it states "
        "what one would have to achieve."
        % (mc["sec_total"]["median"], mc["reagent_opex"]["median"], mc["lcow_net"]["median"],
           closure, ceil)),
    (N, "A point-by-point response to every reviewer comment is provided as a separate file. "
        "The Highlights have been rewritten because the original ones quoted results that the "
        "revision withdraws."),
    (N, "The model implementation, parameter set, random seed and the scripts that generate every "
        "table and figure are openly available at "
        "https://github.com/sandlerleon/iledbv-desalination-model (release v2.3.0, MIT license) "
        "and archived at https://doi.org/10.5281/zenodo.22178234. The manuscript and response to "
        "reviewers are archived at https://doi.org/10.5281/zenodo.22178232."),
    (N, "The manuscript is original, has not been published elsewhere, and is not under "
        "consideration at any other journal. I am the sole author and have approved the revised "
        "manuscript. I have no competing interests to declare, and the work received no external "
        "funding."),
    (N, "Thank you for your consideration."),
    (N, "Sincerely,"),
    (N, "Leon Sandler"),
    (N, "Independent Researcher"),
    (N, "sandler.leon@gmail.com"),
]
rewrite(os.path.join(MS, "Desalination_Cover_Letter_Sandler.docx"),
        os.path.join(MS, "DWT_Revision_Cover_Letter_Sandler.docx"), LETTER)
